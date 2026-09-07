from decimal import Decimal
from pathlib import Path
from datetime import date
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Period, Expense


class ExcelExporter:
    """Exporte une période (figée) en fichier Excel."""

    def __init__(self, output_dir: str = "data/exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_period(self, period_id: int, db: Session) -> Path:
        """
        Exporte une période en Excel.
        Retourne le chemin du fichier créé.
        """
        period = db.get(Period, period_id)
        if not period:
            raise ValueError(f"Période {period_id} non trouvée")

        if period.status != "frozen":
            raise ValueError(f"Période {period_id} n'est pas figée")

        wb = Workbook()
        wb.remove(wb.active)  # Supprime la feuille par défaut

        # Crée les feuilles
        self._add_summary_sheet(wb, period)
        self._add_expenses_sheet(wb, period)
        self._add_by_category_sheet(wb, period)

        # Sauvegarde
        filename = f"periode_{period.id}_{period.start_date.isoformat()}.xlsx"
        filepath = self.output_dir / filename
        wb.save(str(filepath))
        return filepath

    def _add_summary_sheet(self, wb: Workbook, period: Period) -> None:
        """Crée la feuille "Résumé"."""
        ws = wb.create_sheet("Résumé", 0)

        # Styles
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font.color = "FFFFFF"

        title_font = Font(bold=True, size=11)
        total_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

        # Titre
        ws["A1"] = f"Résumé de la période {period.name}"
        ws["A1"].font = title_font
        ws.merge_cells("A1:D1")

        # Infos période
        row = 3
        ws[f"A{row}"] = "Période"
        ws[f"B{row}"] = f"{period.start_date} à {period.end_date}"
        row += 1

        ws[f"A{row}"] = "Statut"
        ws[f"B{row}"] = period.status
        row += 1

        ws[f"A{row}"] = "Dépenses"
        ws[f"B{row}"] = len(period.expenses)
        row += 2

        # Totaux
        ws[f"A{row}"] = "Total général"
        ws[f"B{row}"] = float(period.total_amount or 0)
        ws[f"B{row}"].number_format = '0.00"€"'
        ws[f"A{row}"].fill = total_fill
        ws[f"B{row}"].fill = total_fill
        row += 1

        ws[f"A{row}"] = "Adulte 1"
        ws[f"B{row}"] = float(period.adulte1_total or 0)
        ws[f"B{row}"].number_format = '0.00"€"'
        row += 1

        ws[f"A{row}"] = "Adulte 2"
        ws[f"B{row}"] = float(period.adulte2_total or 0)
        ws[f"B{row}"].number_format = '0.00"€"'

        # Largeurs
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20

    def _add_expenses_sheet(self, wb: Workbook, period: Period) -> None:
        """Crée la feuille "Dépenses"."""
        ws = wb.create_sheet("Dépenses", 1)

        # En-têtes
        headers = ["Date", "Libellé", "Montant", "Catégorie", "Source", "Statut", "Commentaire"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        # Données
        for row_idx, expense in enumerate(period.expenses, start=2):
            ws.cell(row=row_idx, column=1).value = expense.date
            ws.cell(row=row_idx, column=1).number_format = 'yyyy-mm-dd'

            ws.cell(row=row_idx, column=2).value = expense.label
            ws.cell(row=row_idx, column=3).value = float(expense.amount)
            ws.cell(row=row_idx, column=3).number_format = '0.00"€"'

            ws.cell(row=row_idx, column=4).value = expense.category
            ws.cell(row=row_idx, column=5).value = expense.source
            ws.cell(row=row_idx, column=6).value = expense.status
            ws.cell(row=row_idx, column=7).value = expense.comment or ""

        # Largeurs
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 10
        ws.column_dimensions["G"].width = 20

    def _add_by_category_sheet(self, wb: Workbook, period: Period) -> None:
        """Crée la feuille "Par catégorie"."""
        ws = wb.create_sheet("Par catégorie", 2)

        # Agrège par catégorie
        by_cat = {}
        for expense in period.expenses:
            cat = expense.category or "divers"
            if cat not in by_cat:
                by_cat[cat] = Decimal("0")
            by_cat[cat] += expense.amount

        # En-têtes
        ws["A1"] = "Catégorie"
        ws["B1"] = "Montant"
        ws["C1"] = "% Total"

        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["B1"].font = Font(bold=True, color="FFFFFF")
        ws["C1"].font = Font(bold=True, color="FFFFFF")

        ws["A1"].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws["B1"].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws["C1"].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        # Données
        total = period.total_amount or Decimal("0")
        for row_idx, (category, amount) in enumerate(sorted(by_cat.items()), start=2):
            ws.cell(row=row_idx, column=1).value = category
            ws.cell(row=row_idx, column=2).value = float(amount)
            ws.cell(row=row_idx, column=2).number_format = '0.00"€"'

            if total > 0:
                pct = float((amount / total) * 100)
            else:
                pct = 0
            ws.cell(row=row_idx, column=3).value = pct
            ws.cell(row=row_idx, column=3).number_format = '0.0"%"'

        # Largeurs
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
