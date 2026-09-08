#!/usr/bin/env python3
"""Crée un fichier Excel modèle pour importer les repas.

Migré depuis le repo alo standalone le 2026-09-08 (voir mémoire auto
objectif_repo_unique_maisonnettev2). Sortie écrite dans alo-backend/data/exports/,
le dossier de données déjà utilisé par le service (monté en volume Docker).
"""

import sys
from pathlib import Path
from datetime import date, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("❌ openpyxl n'est pas installé. Utilise: pip install openpyxl")
    sys.exit(1)


def create_meals_template():
    """Crée un fichier Excel modèle pour les repas."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Repas"

    # Configuration des colonnes
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12

    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    headers = ["Date", "Compte", "Personne", "Repas"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    # Exemples de remplissage
    accounts = {
        "gourmich": ["Loïc", "Mahaut", "Alban", "Ilan"],
        "tigresse": ["Alice", "Adèle", "Joséphine", "Albert", "Oscar"]
    }

    # Générer des lignes d'exemple (depuis aujourd'hui jusqu'à 7 jours en arrière)
    today = date.today()
    start_date = today - timedelta(days=6)

    row = 2
    for i in range(7):
        current_date = start_date + timedelta(days=i)

        # Alterner entre gourmich et tigresse
        account = "gourmich" if i % 2 == 0 else "tigresse"
        people = accounts[account]

        for person in people:
            cell = ws.cell(row=row, column=1)
            cell.value = current_date
            cell.number_format = 'YYYY-MM-DD'
            cell.alignment = center_align
            cell.border = border

            cell = ws.cell(row=row, column=2)
            cell.value = account
            cell.alignment = center_align
            cell.border = border

            cell = ws.cell(row=row, column=3)
            cell.value = person
            cell.alignment = center_align
            cell.border = border

            cell = ws.cell(row=row, column=4)
            cell.value = 3 if person == "Loïc" or person == "Alice" else 0
            cell.alignment = center_align
            cell.border = border
            cell.number_format = '0'

            row += 1

    # Feuille Instructions
    instructions = wb.create_sheet("Instructions")
    instructions.column_dimensions['A'].width = 100

    instr_text = [
        "INSTRUCTIONS POUR REMPLIR LE FICHIER REPAS",
        "",
        "1. COLONNES À REMPLIR:",
        "   - Date: La date du repas (format YYYY-MM-DD, ex: 2026-07-15)",
        "   - Compte: Le compte (gourmich ou tigresse)",
        "   - Personne: Le nom de la personne",
        "   - Repas: Le nombre de repas (0-4)",
        "",
        "2. COMPTES ET PERSONNES DISPONIBLES:",
        "   Gourmich: Loïc, Mahaut, Alban, Ilan",
        "   Tigresse: Alice, Adèle, Joséphine, Albert, Oscar",
        "",
        "3. FORMAT DE DATE:",
        "   - Utilise le format YYYY-MM-DD (ex: 2026-07-15)",
        "   - Ne remplis qu'un jour par ligne",
        "",
        "4. NOMBRE DE REPAS:",
        "   - Saisis un nombre entre 0 et 4",
        "   - 0 = pas de repas",
        "   - 3 = 3 repas",
        "",
        "5. EXEMPLE:",
        "   Date: 2026-07-15",
        "   Compte: gourmich",
        "   Personne: Loïc",
        "   Repas: 3",
        "",
        "6. APRÈS AVOIR REMPLI LE FICHIER:",
        "   - Exporte-le en format Excel (.xlsx)",
        "   - Va sur le backoffice meals de maisonnettev2",
        "   - Utilise le bouton 'Importer' pour intégrer les données",
        "",
        "NOTES IMPORTANTES:",
        "- N'ajoute pas de nouvelles colonnes",
        "- N'utilise que les noms de personnes de la liste ci-dessus",
        "- La date doit être dans le format spécifié",
    ]

    for row_num, text in enumerate(instr_text, 1):
        cell = instructions.cell(row=row_num, column=1)
        cell.value = text
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if text.startswith("INSTRUCTIONS") or text.startswith("NOTES"):
            cell.font = Font(bold=True, size=12)
        elif text.startswith("   -"):
            cell.font = Font(italic=True)

    # Sauvegarde dans alo-backend/data/exports/ (dossier existant, monté en volume Docker)
    output_path = Path(__file__).parent.parent / "alo-backend" / "data" / "exports" / "meals_template.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"✅ Fichier créé: {output_path}")
    print(f"📊 Exemple de données pré-remplies pour les 7 derniers jours")
    return str(output_path)


if __name__ == "__main__":
    create_meals_template()
